"""Validate post-event CSV/XLSX quality grades, then optionally import them atomically."""
import argparse
import csv
from pathlib import Path
from dotenv import load_dotenv
if __name__=='__main__':
    load_dotenv(Path(__file__).parent/'.env')
from logger_db import init_db
from event_service import EventError,apply_reviews

CRITERIA=('clear_task_0_to_1','relevant_context_0_to_1','constraint_0_to_1','testing_refinement_0_to_1')


def read_reviews(path):
    path=Path(path)
    if path.suffix.lower()=='.csv':
        with path.open(encoding='utf-8-sig',newline='') as source:
            rows=list(csv.DictReader(source))
    elif path.suffix.lower()=='.xlsx':
        from openpyxl import load_workbook
        book=load_workbook(path,read_only=True,data_only=False)
        try:
            if 'QualityReview' not in book:
                raise EventError('The workbook needs a QualityReview worksheet.')
            values=book['QualityReview'].iter_rows(values_only=True)
            headers=next(values,None)
            if not headers:
                raise EventError('The review sheet is empty.')
            rows=[dict(zip(headers,values)) for values in values]
        finally:
            book.close()
    else:
        raise EventError('Choose a .csv or .xlsx review file.')
    result=[]
    for index,row in enumerate(rows,2):
        score=row.get('score')
        criteria=[row.get(name) for name in CRITERIA]
        empty=lambda value: value is None or str(value).strip()==''
        if empty(score) and all(empty(value) for value in criteria):
            continue
        try:
            if isinstance(score,bool) or any(isinstance(value,bool) for value in criteria):
                raise ValueError
            if empty(score):
                if any(empty(value) for value in criteria):
                    raise ValueError
                criteria=[float(value) for value in criteria]
                if not all(0<=value<=1 for value in criteria):
                    raise ValueError
                score=sum(criteria)
            else:
                score=float(score)
        except (ValueError,TypeError):
            raise EventError(f'Invalid score or incomplete rubric in row {index}. Formulas are not accepted; enter numeric grades.')
        result.append({'batch_id':str(row.get('batch_id') or '').strip(),'team_id':str(row.get('team_id') or '').strip(),
            'score':score,'reason':str(row.get('reason') or '').strip()})
    return result


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('file')
    parser.add_argument('--apply',action='store_true',help='Apply all validated rows. The default only validates.')
    args=parser.parse_args()
    try:
        init_db()
        outcome=apply_reviews(read_reviews(args.file),dry_run=not args.apply)
    except EventError as error:
        parser.exit(1,str(error)+'\n')
    except Exception as error:
        parser.exit(1,f'Review import failed ({type(error).__name__}). No grade batch was committed.\n')
    print(f"Validated {outcome['validated']} grades; applied {outcome['applied']}." )
