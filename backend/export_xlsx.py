"""Build a complete per-session judging workbook without exporting access tokens."""
from datetime import datetime, timezone
import io
import json
import re
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from logger_db import connect
import event_service as service

MAX_CELL = 30000


def utc(value):
    if value is None:
        return None
    if isinstance(value,(float,int)):
        return datetime.fromtimestamp(value,timezone.utc).isoformat()
    return str(value)


def text(value):
    if value is None or isinstance(value,(int,float,bool)):
        return value
    value=json.dumps(value,ensure_ascii=False) if isinstance(value,(dict,list)) else str(value)
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]',lambda match: f'\\u{ord(match.group()):04x}',value)


def sheet(book,name,headers,rows):
    ws=book.create_sheet(name)
    lengths={key:max([1]+[(len(str(text(row.get(key,'')) or ''))+MAX_CELL-1)//MAX_CELL for row in rows]) for key in headers}
    expanded=[(key,index) for key in headers for index in range(lengths[key])]
    ws.append([key if lengths[key]==1 else f'{key}_{index+1}' for key,index in expanded])
    for row in rows:
        values=[]
        for key,index in expanded:
            value=text(row.get(key))
            values.append(value[index*MAX_CELL:(index+1)*MAX_CELL] if isinstance(value,str) else value)
        ws.append(values)
        for cell in ws[ws.max_row]:
            if isinstance(cell.value,str):
                cell.data_type='s'
    ws.freeze_panes='A2'
    ws.auto_filter.ref=ws.dimensions
    for cell in ws[1]:
        cell.font=Font(bold=True,color='FFFFFF')
        cell.fill=PatternFill('solid',fgColor='183D56')
        cell.alignment=Alignment(wrap_text=True,vertical='top')
    for index,(key,_) in enumerate(expanded,1):
        ws.column_dimensions[get_column_letter(index)].width=65 if any(term in key for term in ('prompt','response','reason','details','evidence')) else 23
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment=Alignment(wrap_text=True,vertical='top')
    return ws


def export_batch(batch_id=None):
    with connect() as conn:
        # Every sheet uses one database snapshot, including simultaneous grades.
        conn.execute('BEGIN TRANSACTION ISOLATION LEVEL REPEATABLE READ READ ONLY' if conn.postgres else 'BEGIN')
        summary=service.admin_sessions(batch_id,connection=conn)
        batch_id=summary['event']['batch_id']
        teams=[service.team_dict(row,True) for row in conn.execute('SELECT * FROM teams WHERE batch_id=? ORDER BY registered_at',(batch_id,))]
        details=[service.admin_detail(row['session_id'],connection=conn) for row in summary['sessions']]
    book=Workbook()
    book.remove(book.active)
    grade_rows=[]
    for row,detail in zip(summary['sessions'],details):
        row['quality_reason']=detail['session'].get('quality_reason')
        grade_rows.append({'batch_id':batch_id,'team_id':row['team_id'],'team_name':row['team_name'],
            'score':row['quality_score'],'reason':row['quality_reason'],
            'clear_task_0_to_1':None,'relevant_context_0_to_1':None,'constraint_0_to_1':None,'testing_refinement_0_to_1':None})
    sheet(book,'Summary',['batch_id','team_id','team_name','school_name','phase','round1_score','word1_score','word2_score','objective_score','quality_score','quality_status','final_score','prompt_count','quality_reason'],summary['sessions'])
    metadata=[{'field':'batch_id','value':batch_id},{'field':'session','value':summary['event']['name']},
        {'field':'started_at_utc','value':utc(summary['event']['started_at'])},{'field':'ended_at_utc','value':utc(summary['event']['ended_at'])},
        {'field':'exported_at_utc','value':datetime.now(timezone.utc).isoformat()},
        {'field':'scoring','value':'Objective score /6 = timeline /2 + first word /2 + second word /2. Quality /4 is graded after this session ends. Final /10 remains empty until quality is reviewed.'},
        {'field':'rubric','value':'Clear task 0..1; relevant context 0..1; useful constraints or format 0..1; testing or refinement 0..1. Judge semantic quality from the transcript and write evidence in reason. No automatic quality marks.'},
        {'field':'import','value':'Fill QualityReview score and reason, or all four rubric cells and reason. Run grade_reviews.py file.xlsx to validate; add --apply to save all validated rows atomically. Keep batch_id and team_id unchanged.'},
        {'field':'text_storage','value':'All text is stored as literal text. Columns longer than 30000 characters continue in numbered columns. XML control characters are represented with Unicode escape sequences.'},
        {'field':'attempt_history','value':'Append-only history begins with this version. Earlier saved states retain totals and full accepted transcripts, but old individual timeline submissions and word guesses cannot be reconstructed.'}]
    sheet(book,'Event',['field','value'],metadata)
    sheet(book,'Teams',['batch_id','team_id','team_name','school_name','team_size','registered_at','status'],teams)
    members=[{'batch_id':batch_id,'team_id':team['team_id'],'team_name':team['team_name'],'school_name':team['school_name'],**member} for team in teams for member in team['participants']]
    sheet(book,'Participants',['batch_id','team_id','team_name','school_name','name','grade'],members)
    prompts=[]
    timelines=[]
    answers=[]
    audit=[]
    for row,detail in zip(summary['sessions'],details):
        common={'batch_id':batch_id,'team_id':row['team_id'],'team_name':row['team_name'],'session_id':row['session_id']}
        for turn in detail['messages']:
            prompts.append({**common,'request_status':'accepted','turn':turn.get('turn'),'timestamp_utc':utc(turn.get('timestamp')),
                'target':turn.get('target'),'prompt':turn.get('prompt'),'response':turn.get('response'),
                'submitted_at_utc':utc(turn.get('submitted_at')),'response_at_utc':utc(turn.get('response_at')),'latency_seconds':turn.get('latency_seconds'),
                'provider':turn.get('ai_provider',turn.get('provider')),'technique':turn.get('technique',turn.get('detected_technique')),
                'clue_unlocked':turn.get('clue_unlocked'),'quality_status':row['quality_status'],'details':turn})
        for failure in detail['failed_requests']:
            prompts.append({**common,'request_status':failure.get('request_status','failed'),'timestamp_utc':utc(failure['timestamp']),'target':failure.get('target'),
                'prompt':failure['prompt'],'response':failure['error'],'provider':failure['provider_details'].get('ai_provider'),
                'error_reason':failure['ai_fallback_reason'],'retry_after_seconds':failure['ai_retry_after_seconds'],'details':failure['provider_details']})
        for attempt in detail['timeline_attempts']:
            timelines.append({**common,**attempt,'timestamp_utc':utc(attempt['submitted_at'])})
        for answer in detail['word_answers']:
            answers.append({**common,**answer,'timestamp_utc':utc(answer['submitted_at'])})
        for review in detail['quality_audit']:
            audit.append({**common,**review,'reviewed_at_utc':utc(review['reviewed_at'])})
    sheet(book,'Prompts',['batch_id','team_id','team_name','session_id','request_status','turn','timestamp_utc','submitted_at_utc','response_at_utc','latency_seconds','target','prompt','response','provider','technique','clue_unlocked','quality_status','error_reason','retry_after_seconds','details'],prompts)
    sheet(book,'TimelineAttempts',['batch_id','team_id','team_name','session_id','id','timestamp_utc','attempt','placed_ids','slot_results','correct','score'],timelines)
    sheet(book,'WordAnswers',['batch_id','team_id','team_name','session_id','id','timestamp_utc','word_index','answer','correct','outcome'],answers)
    review=sheet(book,'QualityReview',['batch_id','team_id','team_name','score','reason','clear_task_0_to_1','relevant_context_0_to_1','constraint_0_to_1','testing_refinement_0_to_1'],grade_rows)
    validation=DataValidation(type='decimal',operator='between',formula1=0,formula2=4,allow_blank=True)
    validation.error='Enter a score from 0 to 4.'
    validation.showErrorMessage=True
    review.add_data_validation(validation)
    if grade_rows:
        validation.add(f'D2:D{len(grade_rows)+1}')
    sheet(book,'ReviewAudit',['batch_id','team_id','team_name','session_id','score','reason','reviewed_at_utc'],audit)
    output=io.BytesIO()
    book.save(output)
    return output.getvalue(),batch_id
