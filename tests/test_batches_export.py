"""Independent event batches and complete post-event judging exports."""
import io
from contextlib import closing
import json
from pathlib import Path
import sqlite3
import sys
import tempfile
import time
import unittest
from unittest.mock import patch
from openpyxl import load_workbook

sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'backend'))
import event_service as service
import logger_db
from export_xlsx import export_batch
from grade_reviews import read_reviews
from timeglitch_engine import TimeglitchSession

class BatchTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        for change in (patch.object(logger_db,'DATABASE_URL',''),patch.object(logger_db,'DB_PATH',Path(self.temp.name)/'event.sqlite3')):
            change.start()
            self.addCleanup(change.stop)
        logger_db.init_db()

    def register(self,name='Team',batch_id=None):
        return service.register({'teamName':name,'schoolName':'School','batch_id':batch_id,'participants':[{'name':'Student','grade':'11'}]})

    def engine(self,team):
        with logger_db.connect() as conn:
            registered=service.authorized_team(conn,team['team_token'])
            return service.load(conn,registered)[1]

    def solve_r1(self,team):
        engine=self.engine(team)
        return service.mutate(team['team_token'],'timeline',{'placed_ids':engine.correct_order})

    def test_morning_and_afternoon_remain_independent(self):
        morning=self.register('Same team','morning')
        first=service.start_event('morning')
        self.solve_r1(morning)
        with self.assertRaises(service.EventError):
            service.activate_batch('afternoon')
        service.end_event('morning')
        service.activate_batch('afternoon')
        with self.assertRaises(service.EventError):
            self.register('Stale form','morning')
        afternoon=self.register('SAME TEAM','afternoon')
        self.assertNotEqual(morning['team']['team_id'],afternoon['team']['team_id'])
        second=service.start_event('afternoon')
        self.assertGreater(second['started_at'],first['started_at'])
        old=service.get_state(morning['team_token'])
        new=service.get_state(afternoon['team_token'])
        self.assertEqual((old['phase'],old['batch_id']),('finished','morning'))
        self.assertEqual((new['phase'],new['batch_id']),('round1','afternoon'))
        self.assertEqual(service.admin_sessions('morning')['total'],1)
        self.assertEqual(service.admin_sessions('afternoon')['total'],1)
        service.end_event('afternoon')
        self.assertEqual(service.get_state(morning['team_token'])['objective_score'],2)

    def test_postevent_review_batch_validation_and_atomic_import(self):
        one,two=self.register('One'),self.register('Two')
        engine=self.engine(one)
        with self.assertRaises(service.EventError) as caught:
            service.review_quality(engine.session_id,{'score':4,'reason':'Too early.'})
        self.assertEqual(caught.exception.status,409)
        service.end_event('morning')
        valid={'batch_id':'morning','team_id':one['team']['team_id'],'score':3.5,'reason':'Evidence from transcript.'}
        self.assertEqual(service.apply_reviews([valid])['applied'],0)
        self.assertIsNone(service.get_state(one['team_token'])['final_score'])
        with self.assertRaises(service.EventError):
            service.apply_reviews([valid,{'batch_id':'afternoon','team_id':two['team']['team_id'],'score':4,'reason':'Wrong batch.'}],dry_run=False)
        with logger_db.connect() as conn:
            self.assertEqual(conn.execute('SELECT COUNT(*) FROM quality_audit').fetchone()[0],0)
        self.assertEqual(service.apply_reviews([valid],dry_run=False)['applied'],1)
        self.assertEqual(service.get_state(one['team_token'])['final_score'],3.5)
        with self.assertRaises(service.EventError):
            service.apply_reviews([valid,valid],dry_run=False)

    def test_full_workbook_and_failed_request_telemetry(self):
        team=self.register('=SUM(1,1)')
        service.start_event()
        engine=self.engine(team)
        service.mutate(team['team_token'],'timeline',{'placed_ids':list(reversed(engine.correct_order))})
        self.solve_r1(team)
        prompt='=SUM(1,1) Explain a cipher hint for the first word using an unrelated example.'
        def unavailable(*args,**kwargs):
            kwargs['metadata'].update(ai_fallback_reason='org_cooldown',ai_retry_after_seconds=30)
            return '',True
        with patch('timeglitch_engine.get_chronos_response',side_effect=unavailable):
            with self.assertRaises(service.EventError) as caught:
                service.prompt(team['team_token'],{'message':prompt,'target':1})
            self.assertEqual(caught.exception.status,503)
            self.assertTrue(caught.exception.details['retryable'])
        self.assertEqual(service.get_state(team['team_token'])['round2']['prompt_count'],0)
        with logger_db.transaction() as conn:
            conn.execute('UPDATE game_sessions SET last_prompt_at=0 WHERE session_id=?',(engine.session_id,))
        full_response='=CMD('+('R'*70000)
        with patch('timeglitch_engine.get_chronos_response',return_value=(full_response,False)):
            service.prompt(team['team_token'],{'message':prompt,'target':1})
        service.mutate(team['team_token'],'answer',{'word_index':1,'answer':'WRONG'})
        for index,word in enumerate(engine.challenge['words'],1):
            service.mutate(team['team_token'],'answer',{'word_index':index,'answer':word['answer']})
        service.mutate(team['team_token'],'answer',{'word_index':2,'answer':'REPLAY'})
        service.end_event()
        payload,batch=export_batch('morning')
        book=load_workbook(io.BytesIO(payload),data_only=False)
        self.assertEqual(batch,'morning')
        self.assertEqual(set(book.sheetnames),{'Summary','Event','Teams','Participants','Prompts','TimelineAttempts','WordAnswers','QualityReview','ReviewAudit'})
        self.assertEqual(book['TimelineAttempts'].max_row,3)
        self.assertEqual(book['WordAnswers'].max_row,5)
        self.assertEqual(book['Prompts'].max_row,3)
        headers=[cell.value for cell in book['Prompts'][1]]
        rows=list(book['Prompts'].iter_rows(min_row=2,values_only=True))
        accepted=next(row for row in rows if row[headers.index('request_status')]=='accepted')
        reconstructed=''.join(accepted[index] or '' for index,name in enumerate(headers) if name.startswith('response_') and name[9:].isdigit())
        self.assertEqual(reconstructed,full_response)
        for ws in book:
            for row in ws:
                for cell in row:
                    self.assertNotEqual(cell.data_type,'f')
                    self.assertLessEqual(len(str(cell.value or '')),30000)
        review=book['QualityReview']
        review['D2']=4
        review['E2']='Judges checked the transcript against all four criteria.'
        path=Path(self.temp.name)/'grades.xlsx'
        book.save(path)
        imported=read_reviews(path)
        self.assertEqual(service.apply_reviews(imported,dry_run=False)['applied'],1)
        final=service.get_state(team['team_token'])
        self.assertEqual(final['objective_score'],6)
        self.assertEqual(final['final_score'],10)

    def test_old_sqlite_global_name_constraint_migrates_safely(self):
        source=Path(self.temp.name)/'legacy.sqlite3'
        with closing(sqlite3.connect(source)) as conn:
            conn.execute('CREATE TABLE event(id INTEGER PRIMARY KEY,started_at REAL,ended_at REAL)')
            conn.execute('INSERT INTO event(id) VALUES(1)')
            conn.execute("CREATE TABLE teams(team_id TEXT PRIMARY KEY,team_name TEXT NOT NULL COLLATE NOCASE UNIQUE,school_name TEXT NOT NULL,participants TEXT NOT NULL,registered_at REAL NOT NULL,token_hash TEXT NOT NULL UNIQUE,status TEXT NOT NULL DEFAULT 'Confirmed')")
            conn.execute('INSERT INTO teams(team_id,team_name,school_name,participants,registered_at,token_hash) VALUES(?,?,?,?,?,?)',('old','Legacy','School','[]',time.time(),'old-token-hash'))
            conn.commit()
        with patch.object(logger_db,'DB_PATH',source):
            logger_db.init_db()
            with logger_db.connect() as conn:
                self.assertEqual(conn.execute('SELECT batch_id FROM teams WHERE team_id=?',('old',)).fetchone()[0],'morning')
                conn.execute("INSERT INTO teams(team_id,team_name,school_name,participants,registered_at,token_hash,batch_id) VALUES('new','LEGACY','School','[]',1,'new-token-hash','afternoon')")
                self.assertEqual(conn.execute('SELECT COUNT(*) FROM teams').fetchone()[0],2)

if __name__=='__main__':
    unittest.main(verbosity=2)
