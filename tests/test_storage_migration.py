"""Storage contracts and optional real PostgreSQL migration checks."""
from contextlib import contextmanager
import json
import os
from pathlib import Path
import secrets
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'backend'))
import logger_db
import event_service as service
from migrate_sqlite import migrate
from timeglitch_engine import TimeglitchSession

ADMIN_URL = os.getenv('PG_ADMIN_TEST_URL', '')


class StorageContractTests(unittest.TestCase):
    def test_wal_initialization_retries_only_sqlite_lock_errors(self):
        import sqlite3
        from unittest.mock import Mock
        for code in (sqlite3.SQLITE_BUSY,sqlite3.SQLITE_LOCKED,sqlite3.SQLITE_BUSY | (1<<8)):
            with self.subTest(code=code):
                connection=Mock()
                locked=sqlite3.OperationalError('database is locked')
                locked.sqlite_errorcode=code
                attempts=0
                def execute(query):
                    nonlocal attempts
                    if query=='PRAGMA journal_mode=WAL':
                        attempts+=1
                        if attempts==1:
                            raise locked
                    return Mock()
                connection.execute.side_effect=execute
                with patch.object(logger_db.time,'sleep') as pause:
                    logger_db.enable_sqlite_wal(connection)
                self.assertEqual(attempts,2)
                pause.assert_called_once()
                connection.execute.assert_any_call('PRAGMA busy_timeout=15000')

    def test_wal_initialization_does_not_retry_other_errors_or_expired_budget(self):
        import sqlite3
        from unittest.mock import Mock
        for code in (sqlite3.SQLITE_READONLY,sqlite3.SQLITE_CORRUPT,None,sqlite3.SQLITE_BUSY):
            with self.subTest(code=code):
                connection=Mock()
                failure=sqlite3.OperationalError('initialization failed')
                if code is not None:
                    failure.sqlite_errorcode=code
                def execute(query):
                    if query=='PRAGMA journal_mode=WAL':
                        raise failure
                    return Mock()
                connection.execute.side_effect=execute
                with patch.object(logger_db.time,'sleep') as pause:
                    with self.assertRaises(sqlite3.OperationalError) as caught:
                        logger_db.enable_sqlite_wal(connection,timeout=0 if code==sqlite3.SQLITE_BUSY else 5)
                self.assertIs(caught.exception,failure)
                pause.assert_not_called()
                self.assertEqual(sum(call.args[0]=='PRAGMA journal_mode=WAL' for call in connection.execute.call_args_list),1)
                connection.execute.assert_any_call('PRAGMA busy_timeout=15000')

    def test_hosting_requires_persistent_database(self):
        with patch.object(logger_db, 'DATABASE_URL', ''):
            for flag in ({'RENDER':'true'}, {'RENDER_SERVICE_ID':'test-service'}, {'REQUIRE_DATABASE_URL':'true'}):
                with patch.dict(os.environ, flag):
                    with self.assertRaises(RuntimeError):
                        logger_db.connect()

    def test_sqlite_incremental_history_contract(self):
        with tempfile.TemporaryDirectory() as directory:
            with patch.object(logger_db, 'DATABASE_URL', ''), patch.object(logger_db, 'DB_PATH', Path(directory)/'event.sqlite3'):
                logger_db.init_db()
                registered = service.register({'teamName':'Test','schoolName':'School','participants':[{'name':'Student','grade':'11'}]})
                service.start_event()
                with logger_db.transaction() as conn:
                    team = service.authorized_team(conn, registered['team_token'])
                    _, engine = service.load(conn, team)
                    engine.check_timeline_submission(engine.correct_order)
                    engine.prompt_history = [{'turn':1,'prompt':'First','response':'First reply'}, {'turn':2,'prompt':'Second','response':'Second reply'}]
                    engine.round2_prompt_count = 2
                    service.save(conn, engine)
                full = service.get_state(registered['team_token'])
                self.assertTrue(full['prompt_history_complete'])
                self.assertEqual(len(full['prompt_history']), 2)
                delta = service.get_state(registered['team_token'], history_after=1)
                self.assertFalse(delta['prompt_history_complete'])
                self.assertEqual([turn['turn'] for turn in delta['prompt_history']], [2])
                empty = service.get_state(registered['team_token'], history_after=2)
                self.assertEqual(empty['prompt_history'], [])


@unittest.skipUnless(ADMIN_URL, 'Set PG_ADMIN_TEST_URL for isolated local PostgreSQL migration tests.')
class PostgresMigrationTests(unittest.TestCase):
    def setUp(self):
        import psycopg
        from psycopg import sql
        from psycopg.conninfo import conninfo_to_dict, make_conninfo
        if conninfo_to_dict(ADMIN_URL).get('host') not in ('127.0.0.1','localhost','::1'):
            self.fail('These tests only create databases on a local PostgreSQL server.')
        self.temp = tempfile.TemporaryDirectory()
        self.source = Path(self.temp.name)/'source.sqlite3'
        self.old_url, self.old_path = logger_db.DATABASE_URL, logger_db.DB_PATH
        self.database = 'aarohan_test_migration_' + secrets.token_hex(6)
        with psycopg.connect(ADMIN_URL, autocommit=True, connect_timeout=5) as conn:
            conn.execute(sql.SQL('CREATE DATABASE {}').format(sql.Identifier(self.database)))
        self.url = make_conninfo(ADMIN_URL, dbname=self.database)
        logger_db.close_pool()
        logger_db.DATABASE_URL, logger_db.DB_PATH = '', self.source
        logger_db.init_db()
        self.team = service.register({'teamName':'Migration Team','schoolName':'School','participants':[{'name':'Student','grade':'11'}]})
        service.start_event()
        with logger_db.transaction() as conn:
            team = service.authorized_team(conn, self.team['team_token'])
            _, engine = service.load(conn, team)
            engine.check_timeline_submission(engine.correct_order)
            engine.submit_word(1, engine.challenge['words'][0]['answer'])
            engine.prompt_history = [{'turn':1,'prompt':'First','response':'First reply'}, {'turn':2,'prompt':'Second','response':'Second reply'}]
            engine.round2_prompt_count = 2
            service.save(conn, engine)
        self.session_id = engine.session_id
        service.end_event()
        service.review_quality(self.session_id, {'score':3.5,'reason':'Migration review fixture.'})
        logger_db.DATABASE_URL = self.url

    def tearDown(self):
        import psycopg
        from psycopg import sql
        logger_db.close_pool()
        logger_db.DATABASE_URL, logger_db.DB_PATH = self.old_url, self.old_path
        if not self.database.startswith('aarohan_test_migration_'):
            raise AssertionError('Unexpected database name.')
        with psycopg.connect(ADMIN_URL, autocommit=True, connect_timeout=5) as conn:
            conn.execute(sql.SQL('DROP DATABASE {} WITH (FORCE)').format(sql.Identifier(self.database)))
        self.temp.cleanup()

    def test_preserves_state_access_and_audits_without_overwrite(self):
        counts = migrate(self.source)
        self.assertEqual(counts['teams'], 1)
        self.assertEqual(counts['quality_audit'], 1)
        state = service.get_state(self.team['team_token'])
        self.assertEqual(state['total_score'], 7.5)
        self.assertEqual(state['phase'], 'finished')
        self.assertEqual(len(state['prompt_history']), 2)
        self.assertEqual(service.admin_detail(self.session_id)['quality_audit'][0]['score'], 3.5)
        with self.assertRaises(ValueError):
            migrate(self.source)
        service.review_quality(self.session_id, {'score':4,'reason':'New review after migration.'})
        with logger_db.connect() as conn:
            self.assertEqual(conn.execute('SELECT MAX(id) FROM quality_audit').fetchone()[0], 2)
            self.assertEqual(conn.execute('SELECT COUNT(*) FROM teams').fetchone()[0], 1)

    def test_postgres_filters_history_before_network_transfer(self):
        migrate(self.source)
        queries = []
        original = logger_db.Connection.execute
        def capture(instance, query, parameters=None):
            queries.append(query)
            return original(instance, query, parameters)
        with patch.object(logger_db.Connection, 'execute', capture):
            delta = service.get_state(self.team['team_token'], history_after=1)
            summary = service.admin_sessions()
        self.assertEqual([turn['turn'] for turn in delta['prompt_history']], [2])
        self.assertEqual(summary['total'], 1)
        self.assertTrue(any("state::jsonb - 'prompt_history'" in query for query in queries))
        self.assertTrue(any('jsonb_array_elements' in query for query in queries))
        self.assertFalse(any('SELECT * FROM game_sessions' in query for query in queries))
        full = service.get_state(self.team['team_token'])
        self.assertEqual(len(full['prompt_history']), 2)
        from export_xlsx import export_batch
        import io
        from openpyxl import load_workbook
        exported,batch_id=export_batch('morning')
        self.assertEqual(batch_id,'morning')
        workbook=load_workbook(io.BytesIO(exported))
        self.assertEqual(workbook['Prompts'].max_row,3)
        service.end_event()
        self.assertEqual(len(service.admin_detail(self.session_id)['messages']), 2)

    def test_worker_schema_check_does_not_block_live_game_readers(self):
        import concurrent.futures
        migrate(self.source)
        reader=logger_db.connect()
        reader.execute('SELECT * FROM teams').fetchall()
        reader.execute('SELECT * FROM game_sessions').fetchall()
        with concurrent.futures.ThreadPoolExecutor(max_workers=1) as executor:
            check=executor.submit(logger_db.init_db)
            try:
                # A schema check during another worker's active read must not
                # request ALTER TABLE locks on already migrated game tables.
                check.result(timeout=2)
            finally:
                reader.close()


if __name__ == '__main__':
    unittest.main(verbosity=2)
